import jieba
from app.utils.helper import EmailSpider, Saver, EmailSender, get_smtp_host, create_dir
from celery import Celery
import docx
import os
from . import redis_conn, attachment_dir
from openpyxl import Workbook

celery = Celery('tasks', broker='redis://localhost:6379/0', backend='redis://localhost:6379/1')


# todo 检测单用户多账号同时操作，并提供防混乱机制
@celery.task
def download_email(email, email_password, start_time, end_time, report_name):
    saver = Saver(email, email_password, start_time, end_time, report_name)
    spider = EmailSpider(saver)
    attachments, sender_emails, save_dir = spider.run()
    print('附件列表：', attachments)
    print('emails:', sender_emails)

    # for i in sender_emails:
    if sender_emails:
        redis_conn.sadd(email + 'email_set', *sender_emails)
    save_attachments_and_emails(save_dir, email)


# todo 查阅redis连接是否需要释放问题
def save_attachments_and_emails(save_dir, email):
    redis_conn.set('save_dir' + email, save_dir)


@celery.task
def send_receipt(email, email_password, receive_list, subject, content):
    host = get_smtp_host(email)
    outbox = EmailSender()
    outbox.conn_server(email, email_password, host)
    outbox.send_email(receive_list, subject, content)


@celery.task
def jaccard(file_dir, email):
    print("file_dir: " + file_dir)
    text_list = save_text(file_dir)
    result = dict()  # key = filename1, value = [filename2, jaccard_score]
    files = os.listdir(file_dir)
    # TODO 在计算概率的时候只求指定类型的文件
    for file in files:  # 初始化
        result[file] = ['', 0]
    length = len(text_list)

    for index, i in enumerate(text_list):
        for j in text_list:
            if i[0] != j[0]:  # i与j不是同一个文件
                ret1 = i[1].intersection(j[1])
                ret2 = i[1].union(j[1])
                jaccard_score = 1.0 * len(ret1) / len(ret2)
                jaccard_score = round(jaccard_score, 3)  # 保留3位小数
                if result[i[0]][1] < jaccard_score:
                    result[i[0]][0] = j[0]
                    result[i[0]][1] = jaccard_score
                if result[j[0]][1] < jaccard_score:
                    result[j[0]][0] = i[0]
                    result[i[0]][1] = jaccard_score
        redis_conn.set(email + 'dump', '{:.2%}'.format(index/length))
    name = email + ':duplicate_result'
    print('查重结果：', result)
    for k, v in result.items():
        tmp = [str(arg) for arg in v]
        value = ','.join(tmp)
        redis_conn.hset(name, k, value)


def cut_text(text):
    words = set()
    for word in jieba.lcut_for_search(text):
        words.add(word)
    return words


def save_text(file_dir):
    files = os.listdir(file_dir)
    text_list = []
    for file in files:
        path = '{0}/{1}'.format(file_dir, file)
        text = get_text(path)
        if text:
            text_set = cut_text(text)
            tmp = [file, text_set]
            text_list.append(tmp)
        else:
            continue
    return text_list


def get_text(path):
    print("path：" + path)
    if path.endswith('docx'):
        file = docx.Document(path)  # 也就是说这里只能解析docx格式的文件
        t = [para.text for para in file.paragraphs]
        text = ''.join(t)
        return text
    else:
        return None

@celery.task
def send_remail(email):
    receipt = redis_conn.hgetall(email + ':receipt')
    redis_conn.delete(email + ':receipt')  # 用完了就删除
    if receipt:
        print('receipt', receipt)
        context = receipt.get(b'on')
        print('context', context)
        if context:
            key = email + 'email_set'
            email_set = redis_conn.smembers(key)
            # redis_conn.delete(key)
            email_password = redis_conn.get(email + '_password')
            redis_conn.delete(key, email + '_password')
            print('email_password', email_password)

            email_sender = EmailSender()
            if email_password and email_set:
                email_list = [e.decode('utf8') for e in email_set]
                print('email_list', email_list)
                connect = email_sender.conn_server(account=email, password=email_password.decode('utf8'),
                                                   host=get_smtp_host(email))
                if connect:
                    context = context.decode('utf8')
                    state = email_sender.send_email(email_list, '作业提交成功通知',
                                                    context, )
                    print("邮件发送成功！")


@celery.task
def generate_report(email):
    name = email + ':duplicate_result'
    res = redis_conn.hgetall(name)
    redis_conn.delete(name)
    result = []

    for k, v in res.items():
        new_k = k.decode('utf-8')
        new_v = v.decode('utf-8')
        splited_v = new_v.rsplit(',', 1)
        result.append([new_k, splited_v[0], splited_v[1]])
    result.sort(key=lambda x: x[2], reverse=True)
    result.insert(0, ['文件名', '与它最相似的文件', '相似度'])
    export_excel(result, email)


def export_excel(data, email):
    wb = Workbook()
    ws = wb.active
    print("Excel_content", data)
    for i in range(len(data)):
        for j in range(len(data[0])):
            c = ws.cell(row=i + 1, column=j + 1)
            c.value = data[i][j]

    xls_path = redis_conn.get(email + 'xls_path')
    if xls_path:
        path = '{0}/{1}/{2}/{3}'.format(attachment_dir, email, 'Xlsxs', xls_path.decode('utf8'))
        save_path = create_dir(path) + '.xlsx'
        os.rmdir(path)
        wb.save(save_path)
